from io import BytesIO
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, Reference
from math import ceil
from openpyxl.chart.label import DataLabelList, DataLabel
from openpyxl.chart.series import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from collections import Counter
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET


PROFILE_KINDS = [
    ('empleado', 'Empleados', '2563EB'),
    ('empresa', 'Empresas', '047857'),
    ('estudio', 'Estudios', '7C3AED'),
    ('sindicato', 'Sindicatos', 'B45309'),
]


def profile_distribution(users):
    counts = Counter(user.get('account_type') for user in users)
    parts = [(key, label, color, counts[key]) for key, label, color in PROFILE_KINDS]
    other = sum(count for key, count in counts.items() if key not in {part[0] for part in PROFILE_KINDS})
    if other:
        parts.append(('other', 'Sin clasificar', '64748B', other))
    return parts


def add_profile_chart(summary, users):
    parts = profile_distribution(users)
    total_row = 15 + len(parts)
    summary['A13'] = 'Tipos de usuarios'
    summary['A13'].font = Font(bold=True, size=13, color='15335A')
    for column, text in enumerate(['Tipo de usuario', 'Cantidad', 'Porcentaje'], 1):
        cell = summary.cell(14, column, text)
        cell.fill = PatternFill('solid', fgColor='15335A')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for row, (_, label, color, count) in enumerate(parts, 15):
        summary.cell(row, 1, label).font = Font(color=color, bold=True)
        summary.cell(row, 2, count).number_format = '#,##0'
        summary.cell(row, 3, f'=IF($B${total_row}=0,0,B{row}/$B${total_row})').number_format = '0.0%'
        for column in (2, 3):
            summary.cell(row, column).alignment = Alignment(horizontal='right', vertical='center')
    summary.cell(total_row, 1, 'Total de cuentas')
    summary.cell(total_row, 2, f'=SUM(B15:B{total_row-1})').number_format = '#,##0'
    summary.cell(total_row, 3, f'=IF(B{total_row}=0,0,SUM(C15:C{total_row-1}))').number_format = '0.0%'
    for cell in summary[total_row][:3]:
        cell.font = Font(bold=True, color='15335A')
        cell.fill = PatternFill('solid', fgColor='EDF3FB')
    summary.cell(total_row+1, 1, 'Base: cuentas filtradas.')
    summary.cell(total_row+1, 2, 'Incluye cuentas sin actividad.')
    summary.cell(total_row+2, 1, 'Fechas: filtran cálculos.')
    summary.cell(total_row+2, 2, 'Porcentajes redondeados a 1 decimal.')
    for row in range(13, total_row+3):
        summary.row_dimensions[row].height = 24
    summary.column_dimensions['B'].width = 28
    summary.column_dimensions['C'].width = 18
    summary.sheet_view.showGridLines = False
    chart_row = total_row + 4
    if not users:
        summary.cell(chart_row, 1, 'No hay cuentas para estos filtros.')
        return
    # El título y la leyenda están en celdas, fuera del objeto gráfico.
    # Así Excel no puede moverlos sobre el círculo al recalcular su diseño.
    summary.merge_cells(start_row=chart_row, start_column=1, end_row=chart_row, end_column=4)
    heading = summary.cell(chart_row, 1, 'Tipos de usuarios')
    heading.font = Font(bold=True, size=14, color='15335A')
    heading.alignment = Alignment(horizontal='center', vertical='center')
    summary.row_dimensions[chart_row].height = 28
    pie_row = chart_row + 2
    chart = PieChart()
    chart.title = None
    chart.legend = None
    chart.layout = None
    chart.width = 20
    chart.height = 10
    chart.firstSliceAng = 270
    chart.add_data(Reference(summary, min_col=2, min_row=14, max_row=total_row-1), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=15, max_row=total_row-1))
    chart.dataLabels = DataLabelList(showPercent=True, showVal=False, showCatName=True, showSerName=False, showLegendKey=False, numFmt='0.0%', separator='\n', dLblPos='ctr')
    chart.dataLabels.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1100, b=True, solidFill='FFFFFF')), endParaRPr=CharacterProperties(lang='es-AR'))])
    for index, (_, _, color, count) in enumerate(parts):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = color
        point.graphicalProperties.line.solidFill = 'FFFFFF'
        chart.series[0].data_points.append(point)
        if count == 0:
            chart.dataLabels.dLbl.append(DataLabel(idx=index, showPercent=False, showVal=False, showCatName=False))
    summary.add_chart(chart, f'A{pie_row}')
    chart_rows = ceil(chart.height / 2.54 * 72 / 24)
    for row in range(pie_row, pie_row + chart_rows + 1):
        summary.row_dimensions[row].height = 24
    legend_row = pie_row + chart_rows + 1
    for index, (_, label, color, _) in enumerate(parts):
        # Cuatro columnas; un eventual perfil antiguo usa una segunda fila.
        row, column = legend_row + index // 4, index % 4 + 1
        cell = summary.cell(row, column, '■ ' + label)
        cell.font = Font(bold=True, size=11, color=color)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        summary.row_dimensions[row].height = 26
    wb = summary.parent
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = 'auto'



def build_usage_xlsx(report):
    wb = Workbook()
    summary = wb.active
    summary.title = 'Resumen'
    summary.append(['Registros y estadísticas de la calculadora'])
    summary.append(['Desde', report['from'], 'Hasta', report['to']])
    summary.append(['Fuente', 'Registro de uso de la calculadora'])
    summary.append(['Cobertura', 'Solo actividad registrada desde la implementación del seguimiento'])
    summary.append(['Zona horaria', 'Argentina (UTC-3)'])
    summary.append(['Perfil', report.get('type') or 'Todos', 'Búsqueda', report.get('search') or ''])
    summary.append([])
    totals = report['totals']
    for title, key in [('Cuentas registradas', 'users'), ('Usuarios con actividad', 'active_users'), ('Cálculos mensuales', 'monthly_count'), ('Cálculos finales', 'final_count')]:
        summary.append([title, int(totals[key])])
    summary.column_dimensions['A'].width = 38
    summary.column_dimensions['B'].width = 65
    summary.column_dimensions['C'].width = 20
    summary.column_dimensions['D'].width = 35
    sheet = wb.create_sheet('Usuarios')
    sheet.append(['Nombre', 'Email', 'Perfil', 'Organización', 'Estado', 'Registro (Argentina)', 'Cálculos mensuales', 'Cálculos finales', 'Último cálculo (Argentina)', 'Novedades'])
    labels = {'empleado': 'Empleado', 'empresa': 'Empresa', 'estudio': 'Estudio contable', 'sindicato': 'Sindicato'}
    def local(value):
        if not value:
            return ''
        return datetime.fromisoformat(value).replace(tzinfo=timezone.utc).astimezone(timezone(timedelta(hours=-3))).strftime('%d/%m/%Y %H:%M')
    for user in report['users']:
        sheet.append([user['name'], user['email'], labels.get(user['account_type'], user['account_type']), user.get('organization_name', ''), user.get('status_label') or ('Dada de baja' if user['status']=='blocked' and str(user.get('deactivated_by_user',False)).lower() in ('true','1') else {'active':'Activa','blocked':'Bloqueada','pending':'Pendiente','inactive':'Inactiva'}.get(user['status'],user['status'])), local(user['created_at']), *[int(user[k]) for k in ['monthly_count','final_count']], local(user.get('last_activity')), 'Aceptadas' if int(user['newsletter_opt_in']) else 'No aceptadas'])
    widths = [30, 40, 22, 35, 14, 23, 22, 22, 25, 20]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    for ws in wb:
        for row in ws:
            for cell in row:
                # Los textos de usuarios nunca se interpretan como fórmulas.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='center')
            ws.row_dimensions[row[0].row].height = 24
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='15335A')
    add_profile_chart(summary, report['users'])
    result = BytesIO()
    wb.save(result)
    # Excel puede heredar el formato entero de la serie si sourceLinked se omite.
    # Esta versión de openpyxl no expone ese atributo en las etiquetas de torta.
    if summary._charts:
        fixed = BytesIO()
        with ZipFile(BytesIO(result.getvalue()), 'r') as source, ZipFile(fixed, 'w', ZIP_DEFLATED) as output:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename.startswith('xl/charts/chart') and entry.filename.endswith('.xml'):
                    tree = ET.fromstring(data)
                    for number_format in tree.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numFmt'):
                        number_format.set('sourceLinked', '0')
                    data = ET.tostring(tree, encoding='utf-8')
                output.writestr(entry, data)
        result = fixed
    result.seek(0)
    return result
